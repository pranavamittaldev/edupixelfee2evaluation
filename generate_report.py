import os
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()
    
    # Title Page
    doc.add_heading('EduPixel Major Capstone Project Report', 0)
    
    # Let's read the markdown file
    md_file = r'd:\newprojectworknamansaini\fee3sem\EduPixel_Major_Capstone_Project_Report.md'
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
        
    lines = content.split('\n')
    
    # Process lines
    in_code_block = False
    for line in lines:
        stripped = line.strip()
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'No Spacing'
            if p.runs:
                p.runs[0].font.name = 'Courier New'
                p.runs[0].font.size = Pt(9)
            continue

        if not stripped:
            continue
            
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
        elif stripped.startswith('!['): # Image
            # e.g. ![Alt](path)
            match = re.search(r'!\[.*?\]\((.*?)\)', stripped)
            if match:
                img_path = match.group(1)
                full_img_path = os.path.join(r'd:\newprojectworknamansaini\fee3sem', img_path)
                if os.path.exists(full_img_path):
                    try:
                        doc.add_picture(full_img_path, width=Inches(6.0))
                    except Exception as e:
                        doc.add_paragraph(f"[Error loading image: {img_path} - {e}]")
                else:
                    doc.add_paragraph(f"[Image not found: {img_path}]")
        else:
            # Paragraph or list
            doc.add_paragraph(stripped)
            
    doc.add_page_break()
    
    doc.add_heading('Code Snippets and Logic', level=1)
    doc.add_paragraph("The following code snippets highlight the core logic used in our application.")
    
    # Insert code snippets
    doc.add_heading('1. Dynamic Role-Based Navigation (auth.js)', level=2)
    code1 = """function renderNavbar() {
    const role = getUserRole();
    const navElement = document.querySelector('header nav');
    const path = window.location.pathname;
    const page = path.substring(path.lastIndexOf('/') + 1) || 'index.html';

    if (role === 'student') {
        navElement.innerHTML = `
            <a href="profile.html">Dashboard</a>
            <a href="attendance.html">Attendance</a>
            <a href="student-fees.html">Fees</a>
            <a href="#" onclick="logout(); return false;">Logout</a>
        `;
    } else if (role === 'admin') {
        navElement.innerHTML = `
            <a href="admin.html">Dashboard</a>
            <a href="admin-classes.html">Attendance</a>
            <a href="admin-search.html">Search</a>
            <a href="#" onclick="logout(); return false;">Logout</a>
        `;
    }
}"""
    p = doc.add_paragraph(code1)
    p.style = 'No Spacing'
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)
    
    doc.add_heading('2. Multi-Filter Student Search (admin-search.html)', level=2)
    code2 = """function applyFilters() {
    const groupVal = document.getElementById('filter-group').value;
    const rollVal = document.getElementById('filter-roll').value.trim().toLowerCase();
    
    rows.forEach(row => {
        const rowGroup = row.getAttribute('data-group');
        const rowRoll = row.getAttribute('data-roll').toLowerCase();
        
        let matchGroup = (groupVal === 'ALL' || rowGroup === groupVal);
        let matchRoll = (!rollVal || rowRoll.includes(rollVal));
        
        if (matchGroup && matchRoll) {
            row.style.display = '';
        } else {
            row.style.display = 'none';
        }
    });
}"""
    p = doc.add_paragraph(code2)
    p.style = 'No Spacing'
    p.runs[0].font.name = 'Courier New'
    p.runs[0].font.size = Pt(9)

    doc.add_page_break()
    
    doc.add_heading('Mobile View Interfaces', level=1)
    doc.add_paragraph("The application is fully responsive. Below are the screenshots of the platform as seen on a mobile device (iPhone 13 viewport).")
    
    mobile_dir = r"d:\newprojectworknamansaini\fee3sem\screenshots\mobile"
    if os.path.exists(mobile_dir):
        mobile_images = os.listdir(mobile_dir)
        for img in mobile_images:
            if img.endswith(".png"):
                doc.add_heading(f'Mobile View: {img.replace("mobile_", "").replace(".png", "")}', level=2)
                try:
                    doc.add_picture(os.path.join(mobile_dir, img), width=Inches(3.5))
                except Exception as e:
                    print(f"Failed to add image {img}: {e}")

    # Save Document
    out_path = r'd:\newprojectworknamansaini\fee3sem\EduPixel_Final_Project_Report.docx'
    doc.save(out_path)
    print(f"Report generated successfully at {out_path}.")

if __name__ == '__main__':
    main()
